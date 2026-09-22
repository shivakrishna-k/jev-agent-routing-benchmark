import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))
import openpyxl

import export_excel


def test_export_runs_without_raw_data(monkeypatch):
    """No API calls: exports whatever exists (here, just the offline baselines) without crashing."""
    with tempfile.TemporaryDirectory() as tmp:
        out_dir = Path(tmp)
        monkeypatch.setattr(sys, "argv", ["export_excel.py", "--dir", str(out_dir)])
        export_excel.main()
        wb = openpyxl.load_workbook(out_dir / "results.xlsx")
        assert wb.sheetnames == ["Read Me", "Cases", "Responses", "Compare", "Summary"]
        assert wb["Cases"].max_row == 101  # header + 100 cases
        # Only the two offline baselines ran (no raw.jsonl), 100 cases each.
        assert wb["Responses"].max_row == 201
